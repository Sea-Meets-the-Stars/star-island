""" Build slides/star_meeting_oct2026.pptx for the October 4, 2026 Star Meeting

Sources: reports/2026-Oct_report.md and slides/eval_tally.json
(written by star_island/survey/tally_evals.py).  All text is >= 20pt.

Usage (from the repo root, in the ocean14 environment):
    python star_island/survey/tally_evals.py
    python slides/build_star_meeting_oct2026.py
"""
import json
import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

here = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(here, 'star_meeting_oct2026.pptx')
with open(os.path.join(here, 'eval_tally.json')) as f:
    tally = json.load(f)

# Palette: deep sea + coral ("brain") accent
INK = '1B2B34'
SEA = '12343B'       # dominant dark
WHITE = 'FFFFFF'
TINT = 'EAF2F0'      # pale sea-glass card fill
CORAL = 'E8735A'
GLASS = '7FB7A4'
GOLD = 'F2C14E'
TEAL = '3E7C8A'
MIST = 'B5C7CC'
PLUM = '8C5A7A'

HEAD = 'Cambria'
BODY = 'Calibri'
MIN = 20  # smallest font size allowed anywhere in the deck

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
TITLE_ONLY = prs.slide_layouts[5]
slide_no = [0]


def rgb(hex_):
    return RGBColor.from_string(hex_)


# ---------- primitives ----------
def new_slide(dark=False):
    s = prs.slides.add_slide(TITLE_ONLY)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(SEA if dark else WHITE)
    slide_no[0] += 1
    return s


def style_run(run, size=MIN, color=INK, bold=False, italic=False, font=BODY):
    f = run.font
    f.size = Pt(max(size, MIN))
    f.color.rgb = rgb(color)
    f.bold = bold
    f.italic = italic
    f.name = font


def set_bullet(p, on=True):
    pPr = p._p.get_or_add_pPr()
    for tag in ('a:buNone', 'a:buChar', 'a:buFont'):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if on:
        pPr.set('marL', str(Inches(0.35)))
        pPr.set('indent', str(-Inches(0.35)))
        buFont = pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'})
        buChar = pPr.makeelement(qn('a:buChar'), {'char': '•'})
        pPr.append(buFont)
        pPr.append(buChar)
    else:
        pPr.append(pPr.makeelement(qn('a:buNone'), {}))


def fill_frame(tf, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """paras: list of dicts {text, size, color, bold, italic, font, bullet, after}
    or list of such dicts as runs under key 'runs'."""
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, 0)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = para.get('runs', [para])
        for r in runs:
            run = p.add_run()
            run.text = r['text']
            style_run(run, r.get('size', MIN), r.get('color', INK), r.get('bold', False),
                      r.get('italic', False), r.get('font', BODY))
        set_bullet(p, para.get('bullet', False))
        if 'after' in para:
            p.space_after = Pt(para['after'])


def text(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    if isinstance(paras, str):
        paras = [dict(text=paras)]
    fill_frame(tb.text_frame, paras, align, anchor)
    return tb


def shape(slide, kind, x, y, w, h, color, name=None):
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(color)
    sh.line.fill.background()
    sh.shadow.inherit = False
    if name:
        sh.name = name
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = 0.06
    return sh


def badge(slide, x, y, d, color, glyph, size=24, glyph_color=WHITE, name=None):
    """The deck's motif: a filled circle carrying a glyph or number."""
    sh = shape(slide, MSO_SHAPE.OVAL, x, y, d, d, color, name)
    fill_frame(sh.text_frame, [dict(text=glyph, size=size, color=glyph_color, bold=True)],
               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    return sh


def title(slide, txt, dark=False, x=0.6, y=0.4, w=12.1, h=1.0, size=36, anchor=MSO_ANCHOR.MIDDLE):
    t = slide.shapes.title
    t.left, t.top, t.width, t.height = Inches(x), Inches(y), Inches(w), Inches(h)
    fill_frame(t.text_frame, [dict(text=txt, size=size, bold=True, font=HEAD,
                                   color=WHITE if dark else SEA)], PP_ALIGN.LEFT, anchor)


def number(slide, dark=False):
    text(slide, 12.2, 6.85, 0.8, 0.45,
         [dict(text=str(slide_no[0]), color=MIST if dark else TEAL)],
         align=PP_ALIGN.RIGHT, name='Slide number')


def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


def card(slide, x, y, w, h, head=None, body=None, fill=TINT, head_color=SEA, body_color=INK,
         name='Card'):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, name + ' card')
    paras = []
    if head:
        paras.append(dict(text=head, size=24, bold=True, color=head_color, after=6))
    if body:
        paras.append(dict(text=body, color=body_color))
    if paras:
        text(slide, x + 0.25, y + 0.2, w - 0.5, h - 0.4, paras, name=name + ' text')


def icon_row(slide, x, y, w, glyph, color, head, body=None, name='Row'):
    badge(slide, x, y, 0.8, color, glyph, name=name + ' badge')
    paras = [dict(text=head, size=22, bold=True, color=SEA)]
    if body:
        paras.append(dict(text=body))
    text(slide, x + 1.05, y - 0.1, w - 1.05, 1.0, paras, anchor=MSO_ANCHOR.MIDDLE,
         name=name + ' text')


def bar_chart(slide, x, y, w, h, labels, values, color, horizontal, name):
    cd = CategoryChartData()
    cd.categories = labels
    cd.add_series('Count', values)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    gf.name = name
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.size = Pt(MIN)
    ch.font.name = BODY
    ch.font.color.rgb = rgb(INK)
    plot = ch.plots[0]
    plot.gap_width = 45
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.font.size = Pt(MIN)
    dl.font.bold = True
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    ser = plot.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = rgb(color)
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    va.minimum_scale = 0
    va.maximum_scale = max(values) * 1.15
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(MIN)
    ca.format.line.color.rgb = rgb(MIST)
    ca.has_major_gridlines = False
    return ch


# ---------- 1. Title ----------
s = new_slide(dark=True)
title(s, 'Big Brain Island', dark=True, x=0.8, y=2.3, w=8.2, h=1.6, size=60,
      anchor=MSO_ANCHOR.BOTTOM)
text(s, 0.8, 4.1, 8.2, 1.9, [
    dict(text="2027 All Star I · Chair's Update", size=26, color=TINT),
    dict(text='Star Meeting · October 4, 2026', size=26, color=GLASS)], name='Subtitle')
shape(s, MSO_SHAPE.OVAL, 9.5, 1.9, 3.0, 3.0, CORAL, 'Title circle')
shape(s, MSO_SHAPE.OVAL, 10.35, 2.75, 1.3, 1.3, SEA, 'Title inner circle')
notes(s, "Chair's update for 2027 All Star I: theme, speaker, and what we heard "
         "in the 2026 evaluations.")

# ---------- 2. Theme ----------
s = new_slide()
title(s, 'Our 2027 theme: the mind, from womb to tomb')
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 1.8, 5.6, 4.7, SEA, 'Theme quote card')
text(s, 0.95, 2.15, 4.9, 4.1, [
    dict(text='Big Brain Island', size=32, bold=True, color=WHITE, font=HEAD, after=12),
    dict(text='Exploring the marvels and the mysteries of the mind from womb to tomb.',
         size=26, italic=True, color=TINT)], name='Theme quote')
text(s, 6.8, 1.8, 5.9, 0.9, [dict(text="Building on 2026's feedback, we want a conference that is:",
                                  size=22)], name='Theme lead-in')
icon_row(s, 6.8, 2.95, 5.9, '1', CORAL, 'Intellectually engaging', name='Engaging')
icon_row(s, 6.8, 4.15, 5.9, '2', TEAL, 'Family-friendly', name='Family')
icon_row(s, 6.8, 5.35, 5.9, '3', PLUM, 'Musically & artistically enriching', name='Music')
number(s)

# ---------- 3. Speaker ----------
s = new_slide()
title(s, 'Speaker: Dr. Caroline Hollnagel')
text(s, 0.6, 1.5, 12.1, 0.9, [dict(
    text='A clinical neuropsychologist who studies how the brain shapes behavior, '
         'and a natural fit for Big Brain Island.', size=22)], name='Speaker summary')
text(s, 0.6, 2.6, 6.4, 0.5, [dict(text='Clinical & research interests', size=24, bold=True,
                                  color=SEA)], name='Interests header')
interests = ['Dementia', 'Frontal lobe & inhibition', 'Stroke & brain injury rehab',
             'First-break schizophrenia', 'Obsessive-compulsive disorder']
for i, t in enumerate(interests):
    col, row = i % 2, i // 2
    w = 6.4 if i == 4 else 3.1
    sh = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6 + col * 3.3, 3.3 + row * 1.05, w, 0.85,
               TINT, f'Interest {i+1}')
    sh.text_frame.margin_left = Inches(0.15)
    fill_frame(sh.text_frame, [dict(text=t, color=SEA)], anchor=MSO_ANCHOR.MIDDLE)
    sh.text_frame.margin_left = Inches(0.15)
text(s, 7.6, 2.6, 5.1, 0.5, [dict(text='Training', size=24, bold=True, color=SEA)],
     name='Training header')
steps = ['Smith College', 'Washington University in St. Louis',
         'UCSF: pre-doctoral internship', 'UCSF Memory & Aging Center: postdoc']
line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(7.85), Inches(3.55),
                              Inches(7.85), Inches(6.1))
line.line.color.rgb = rgb(GLASS)
line.line.width = Pt(3)
for i, t in enumerate(steps):
    y = 3.35 + i * 0.85
    shape(s, MSO_SHAPE.OVAL, 7.65, y, 0.4, 0.4, CORAL, f'Timeline dot {i+1}')
    text(s, 8.3, y - 0.2, 4.4, 0.8, [dict(text=t)], anchor=MSO_ANCHOR.MIDDLE,
         name=f'Timeline step {i+1}')
number(s)
notes(s, 'She has taught and mentored students, residents, and fellows at Washington '
         'University St. Louis, UC Berkeley, and UCSF. Provider profile: '
         'https://doctors.northernlighthealth.org/provider/caroline-hollnagel/1902432')

# ---------- 4. Section: evaluations ----------
s = new_slide(dark=True)
title(s, 'What the 2026 evaluations told us', dark=True, x=0.8, y=2.6, w=7.4, h=1.5, size=44,
      anchor=MSO_ANCHOR.BOTTOM)
text(s, 0.8, 4.3, 7.4, 1.2, [dict(text='This conference is something we all build together.',
                                  size=24, color=TINT)], name='Section subtitle')
shape(s, MSO_SHAPE.OVAL, 8.9, 1.9, 3.4, 3.4, CORAL, 'Responses circle')
text(s, 8.9, 2.2, 3.4, 2.8, [
    dict(text=str(tally['n_responses']), size=88, bold=True, color=WHITE, font=HEAD),
    dict(text='responses', size=26, color=WHITE)], align=PP_ALIGN.CENTER,
     anchor=MSO_ANCHOR.MIDDLE, name='Responses stat')
number(s, dark=True)

# ---------- 5. Favorites chart ----------
s = new_slide()
title(s, 'The Speaker and 15x20 led the favorites')
acts = list(tally['activities'].items())[::-1]  # bar charts plot bottom-up
bar_chart(s, 0.4, 1.5, 8.3, 5.4, [a[0] for a in acts], [a[1] for a in acts], TEAL, True,
          'Favorite activities chart')
n_lec = tally['activities']['Lecture / Speaker']
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 9.0, 1.7, 3.7, 4.7, TINT, 'Lecture callout card')
text(s, 9.3, 2.0, 3.1, 4.1, [
    dict(text=f"{n_lec} of {tally['n_enjoyed']}", size=60, bold=True, color=CORAL, font=HEAD),
    dict(text='named the Lecture or Speaker as a favorite', size=22, color=SEA, after=18),
    dict(text='Counts are respondents who mentioned each activity.', italic=True, color=TEAL)],
     name='Lecture callout')
number(s)
notes(s, 'Tally from star_island/survey/tally_evals.py. Each respondent counts once per '
         f"activity, out of the {tally['n_enjoyed']} who answered 'Which activities did you "
         "enjoy the most?'. Also mentioned: Musicale, karaoke, Book Group, Art Barn, "
         'Club Newton, Pel Show, yoga, kayaking.')

# ---------- 6. Quotes ----------
s = new_slide()
title(s, 'Familiar traditions still delight')
quotes = [('“The speaker was sensational.”', 'On the Lecture', CORAL),
          ('“Morning chapel continues to delight.”', 'Chapel in all its forms', TEAL),
          ('“15 x 20 always just a marvelous way to learn more about people.”', 'On 15x20', PLUM)]
for i, (q, who, col) in enumerate(quotes):
    x = 0.6 + i * 4.15
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.7, 3.85, 3.4, TINT, f'Quote {i+1} card')
    badge(s, x + 0.3, 1.95, 0.75, col, '“', size=36, name=f'Quote {i+1} badge')
    text(s, x + 0.3, 2.9, 3.25, 2.05, [
        dict(text=q, size=24, italic=True, color=SEA, after=10),
        dict(text=who, color=TEAL)], name=f'Quote {i+1} text')
badge(s, 0.6, 5.55, 0.8, GOLD, '★', size=28, glyph_color=SEA, name='Thanks badge')
text(s, 1.65, 5.4, 11.05, 1.1, [dict(
    text="Real gratitude for the Housing and Children's Program coordinators, and for "
         'Cooper and Macy, who kept the Great People Hunt running.')],
     anchor=MSO_ANCHOR.MIDDLE, name='Thanks text')
number(s)

# ---------- 7. Schedule ----------
s = new_slide()
title(s, 'A lighter schedule worked, except on opening night')
for x, head, glyph, col, items in [
        (0.6, 'Keep', '✓', GLASS, [
            dict(text='Fewer things stacked on top of each other'),
            dict(text='“The whole week felt a little less frenetic”', italic=True),
            dict(text="Other meals handled the conference's size fine")]),
        (6.85, 'Fix', '!', CORAL, [
            dict(text='Opening night was overstuffed: marshmallow roast, chapel, and '
                      'get-acquainted all competed'),
            dict(text='Consider moving get-acquainted to Monday'),
            dict(text='Name That Tune moved too fast this year')])]:
    card(s, x, 1.7, 5.85, 4.8, head=head, name=head)
    badge(s, x + 4.85, 1.9, 0.75, col, glyph, size=28, name=head + ' badge')
    for it in items:
        it.update(bullet=True, after=14)
    text(s, x + 0.25, 2.65, 5.35, 3.6, items, name=head + ' list')
number(s)

# ---------- 8. Banquet ----------
s = new_slide()
title(s, "Banquet Night hasn't kept pace with our size")
icon_row(s, 0.6, 1.8, 6.2, '1', CORAL, 'No table inside',
         'Signing up late meant overflow seating', name='Seating')
icon_row(s, 0.6, 3.25, 6.2, '2', CORAL, 'Loud and crowded outside',
         'Hard to hear what was happening inside', name='Noise')
icon_row(s, 0.6, 4.7, 6.2, '3', CORAL, 'Stressful, not special',
         '…and that was in good weather', name='Mood')
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.2, 1.7, 5.5, 4.8, SEA, 'Banquet ideas card')
text(s, 7.5, 1.95, 4.9, 4.3, [
    dict(text='Ideas to fold in', size=24, bold=True, color=GOLD, after=12),
    dict(text="Thank year-round roles: Registrar, Housing, Children's Program, "
              'Executive Committee, Chairs', color=WHITE, bullet=True, after=10),
    dict(text="Include the children's programs", color=WHITE, bullet=True, after=10),
    dict(text='Bring back the centerpiece contest, with decorations made in the Art Barn',
         color=WHITE, bullet=True)], name='Banquet ideas')
number(s)

# ---------- 9. Crowding ----------
s = new_slide()
title(s, 'Crowding is a governance question')
for i, (n, lbl, col) in enumerate([('~300', 'conferees this year', TEAL),
                                    ('10–20', 'fewer people, as one respondent suggested', CORAL)]):
    x = 0.6 + i * 3.4
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.8, 3.1, 3.0, TINT, f'Stat {i+1} card')
    text(s, x + 0.25, 2.0, 2.6, 2.6, [
        dict(text=n, size=60, bold=True, color=col, font=HEAD),
        dict(text=lbl, color=SEA)], name=f'Stat {i+1}')
text(s, 7.6, 1.8, 5.1, 3.0, [
    dict(text='“Too crowded” came up again, both directly and through the Banquet feedback.',
         size=22, after=16),
    dict(text="This is about capacity and finances, so a program tweak alone can't fix it.",
         size=22, bold=True, color=SEA)], name='Crowding text')
badge(s, 0.6, 5.3, 0.8, SEA, '?', size=28, name='Discuss badge')
text(s, 1.65, 5.3, 11.0, 0.8, [dict(text='For the Star Meeting: do we revisit enrollment for 2027?',
                                    size=24, bold=True, color=SEA)],
     anchor=MSO_ANCHOR.MIDDLE, name='Crowding question')
number(s)

# ---------- 10. Ideas ----------
s = new_slide()
title(s, 'New ideas from the evaluations')
ideas = [('Spontaneity hour', 'Unplanned time, announced on the chalkboard'),
         ('Intergenerational', 'More activities that mix ages'),
         ('Teen workshops', 'Adults lead mini-sessions for Jr & Sr Teens'),
         ('Accompanist', 'A dedicated pianist, perhaps grant-funded'),
         ('Inclusion', 'Welcome people beyond the regular cast'),
         ('Fuller intros', 'Introduce Talks on the Rocks speakers properly')]
for i, (h, b) in enumerate(ideas):
    col, row = i % 3, i // 3
    x, y = 0.6 + col * 4.15, 1.7 + row * 2.45
    card(s, x, y, 3.85, 2.15, head=h, body=b, name=f'Idea {i+1}')
    shape(s, MSO_SHAPE.OVAL, x + 3.3, y + 0.28, 0.32, 0.32, [CORAL, TEAL, PLUM][col],
          f'Idea {i+1} dot')
number(s)
notes(s, 'Other suggestions: telescopes for stargazing; reserved served tables for families '
         'with very young kids; junior teens rooming together; longer nestling hours; a '
         "senior-teen MC for the children's stunt show; a breathwork/longevity series; the "
         'Jr Teen boat every other year.')

# ---------- 11. Volunteers ----------
s = new_slide()
title(s, f"{tally['n_volunteer']} people raised a hand to volunteer")
areas = list(tally['volunteer_areas'].items())[::-1]
bar_chart(s, 0.4, 1.5, 7.6, 5.4, [a[0] for a in areas], [a[1] for a in areas], GLASS, True,
          'Volunteer interest chart')
icon_row(s, 8.4, 1.8, 4.3, '✓', TEAL, 'Succession in place', 'Art Barn has two new leaders',
         name='Succession')
icon_row(s, 8.4, 3.3, 4.3, '!', GOLD, 'Recruit earlier', 'Book group needs time to pick books',
         name='Earlier')
icon_row(s, 8.4, 4.8, 4.3, '@', CORAL, 'Check the mailing list',
         'One volunteer missed every call', name='Mailing')
number(s)
notes(s, 'Counts are respondents who checked at least one job in each area. Many also wrote '
         "'happy to help wherever needed'.")

# ---------- 12. Discussion ----------
s = new_slide(dark=True)
title(s, 'For discussion today', dark=True, x=0.8, y=2.6, w=6.8, h=1.5, size=44,
      anchor=MSO_ANCHOR.BOTTOM)
text(s, 0.8, 4.3, 6.8, 1.2, [dict(text='Thank you for building this conference together.',
                                  size=24, color=TINT)], name='Closing subtitle')
for i, t in enumerate(['Re-envision Banquet Night', 'Enrollment and capacity',
                       'A lighter opening night', 'Which new ideas to pilot']):
    y = 1.4 + i * 1.2
    badge(s, 8.0, y, 0.85, CORAL, str(i + 1), size=28, name=f'Ask {i+1} badge')
    text(s, 9.1, y, 3.8, 0.85, [dict(text=t, size=24, color=WHITE)], anchor=MSO_ANCHOR.MIDDLE,
         name=f'Ask {i+1} text')
number(s, dark=True)

prs.core_properties.title = "2027 All Star I Chair's Update"
prs.core_properties.author = 'J. Xavier Prochaska'
prs.save(OUT)
print('Wrote', OUT)
